import io
import pdfplumber
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF binary bytes using pdfplumber."""
    extracted_text = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            text = page.extract_text()
            if text:
                extracted_text.append(text.strip())
    
    return "\n\n".join(extracted_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX binary bytes using python-docx."""
    doc = docx.Document(io.BytesIO(file_bytes))
    extracted_text = []
    
    # Extract text from paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            extracted_text.append(para.text.strip())
            
    # Also extract text from any tables in the resume
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                extracted_text.append(" | ".join(row_text))
                
    return "\n".join(extracted_text)


def parse_resume(file_bytes: bytes, filename: str) -> str:
    """
    Main parser entry point.
    Determines file type and extracts clean text.
    Raises ValueError with user-friendly messages for invalid or unparseable files.
    """
    filename_lower = filename.lower()
    
    if filename_lower.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif filename_lower.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file format for '{filename}'. Only PDF and DOCX are accepted.")

    # Validate that meaningful text was extracted
    if not text or len(text.strip()) < 30:
        raise ValueError(
            "Could not extract sufficient text from the uploaded resume. "
            "The document might be empty, scanned, or contains only images."
        )

    return text.strip()
